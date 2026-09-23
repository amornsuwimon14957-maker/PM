import streamlit as st
import pandas as pd
import io
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# --- 1. การตั้งค่าหน้าจอ Streamlit ---
st.set_page_config(
    page_title="KPI PM Dashboard & Presentation Generator",
    page_icon="⚙️",
    layout="wide"
)

# รายชื่อเครื่องจักรตามไฟล์ KPIs 2024.xlsx
MACHINES = [
    "MDB Tranformer Generator",
    "Tower Crane",
    "Overhead / Gantry Crane",
    "Rebar machine",
    "Hi-steel machine",
    "Batching plant & Flying bucket",
    "Carrousel System"
]

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# ค่าเริ่มต้น Plan/Action รายเดือนตามตัวอย่างข้อมูลจริง
DEFAULT_PLAN = {
    "Jan": [8, 2, 12, 0, 0, 5, 2],
    "Feb": [9, 2, 12, 10, 18, 5, 2],
    "Mar": [7, 2, 14, 0, 14, 5, 2],
    "Apr": [8, 2, 10, 0, 0, 5, 2],
    "May": [7, 3, 11, 10, 0, 5, 2],
    "Jun": [8, 2, 14, 0, 0, 5, 2],
    "Jul": [7, 2, 16, 0, 0, 5, 2],
    "Aug": [8, 2, 10, 10, 18, 5, 2],
    "Sep": [7, 2, 10, 0, 14, 5, 2],
    "Oct": [8, 2, 12, 0, 0, 5, 2],
    "Nov": [7, 3, 15, 10, 0, 5, 2],
    "Dec": [8, 2, 12, 0, 0, 5, 2]
}

DEFAULT_ACTION = {
    "Jan": [8, 2, 12, 0, 0, 5, 2],
    "Feb": [9, 1, 12, 10, 18, 5, 2],
    "Mar": [7, 1, 14, 0, 14, 5, 2],
    "Apr": [8, 2, 10, 0, 0, 5, 2],
    "May": [7, 2, 11, 9, 0, 5, 2],
    "Jun": [8, 1, 14, 0, 0, 5, 2],
    "Jul": [7, 2, 13, 0, 0, 5, 2],
    "Aug": [8, 1, 10, 10, 18, 5, 2],
    "Sep": [7, 1, 10, 0, 14, 5, 2],
    "Oct": [8, 2, 12, 0, 0, 5, 2],
    "Nov": [0, 0, 0, 0, 0, 0, 0],
    "Dec": [0, 0, 0, 0, 0, 0, 0]
}

# --- 2. จัดเก็บข้อมูลใน Session State ---
if "plan_db" not in st.session_state:
    st.session_state.plan_db = pd.DataFrame(DEFAULT_PLAN, index=MACHINES)

if "action_db" not in st.session_state:
    st.session_state.action_db = pd.DataFrame(DEFAULT_ACTION, index=MACHINES)

# --- 3. ส่วน Sidebar ตั้งค่าและเมนูหลัก ---
st.sidebar.title("📌 เมนูและตั้งค่า")
app_mode = st.sidebar.radio("เลือกหน้าการทำงาน", ["📝 กรอกข้อมูลรายเดือน", "📊 Dashboard & สรุปภาพรวม"])

st.sidebar.markdown("---")
year = st.sidebar.text_input("ปีงบประมาณ", "2567")
target_kpi = st.sidebar.number_input("เป้าหมาย KPI (%)", value=85.0, step=1.0)
prepared_by = st.sidebar.text_input("ผู้จัดทำ", "คุณอมร อมรกุล")
approved_by = st.sidebar.text_input("ผู้อนุมัติ", "คุณชาญณรงค์ ตั้งจิตมานะกิจ")


# --- 4. ฟังก์ชันสร้างไฟล์ PowerPoint (.pptx) ---
def generate_pptx_slide(selected_month, plan_series, action_series):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Header
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"KPI การซ่อมบำรุงเชิงป้องกัน (PM) - เดือน {selected_month} {year}"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(24, 43, 73)
    
    p_sub = tf.add_paragraph()
    p_sub.text = f"ผู้จัดทำ: {prepared_by}   |   ผู้อนุมัติ: {approved_by}"
    p_sub.font.size = Pt(14)
    p_sub.font.color.rgb = RGBColor(100, 110, 120)

    # Calculate Totals
    t_plan = int(plan_series.sum())
    t_action = int(action_series.sum())
    kpi_pct = (t_action / t_plan * 100) if t_plan > 0 else 0.0
    is_pass = kpi_pct >= target_kpi

    # Summary Box
    stat_box = slide.shapes.add_shape(1, Inches(0.8), Inches(1.7), Inches(11.733), Inches(1.0))
    stat_box.fill.solid()
    stat_box.fill.fore_color.rgb = RGBColor(240, 244, 248)
    stat_box.line.color.rgb = RGBColor(200, 210, 220)
    
    tf_stat = stat_box.text_frame
    p_s = tf_stat.paragraphs[0]
    status_text = "🟢 ผ่านเกณฑ์" if is_pass else "🔴 ไม่ผ่านเกณฑ์"
    p_s.text = f"เป้าหมาย KPI: > {target_kpi:.0f}%   |   แผนงาน: {t_plan} ครั้ง   |   ทำจริง: {t_action} ครั้ง   |   ผล KPI: {kpi_pct:.2f}% ({status_text})"
    p_s.font.size = Pt(16)
    p_s.font.bold = True
    p_s.font.color.rgb = RGBColor(34, 197, 94) if is_pass else RGBColor(239, 68, 68)

    # Table Data
    rows = len(MACHINES) + 2
    cols = 3
    table_shape = slide.shapes.add_table(rows, cols, Inches(0.8), Inches(3.0), Inches(11.733), Inches(3.8))
    table = table_shape.table
    
    headers = ["รายการเครื่องจักร (Area / Machine)", "แผนงาน (Plan)", "ทำจริง (Action)"]
    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = text
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(30, 41, 59)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.color.rgb = RGBColor(255, 255, 255)
            paragraph.font.size = Pt(13)

    for idx, machine in enumerate(MACHINES):
        table.cell(idx + 1, 0).text = machine
        table.cell(idx + 1, 1).text = str(int(plan_series[machine]))
        table.cell(idx + 1, 2).text = str(int(action_series[machine]))

    # Row สรุปรวม
    table.cell(rows - 1, 0).text = "Total"
    table.cell(rows - 1, 1).text = str(t_plan)
    table.cell(rows - 1, 2).text = str(t_action)
    for c in range(3):
        table.cell(rows - 1, c).fill.solid()
        table.cell(rows - 1, c).fill.fore_color.rgb = RGBColor(226, 232, 240)
        for paragraph in table.cell(rows - 1, c).text_frame.paragraphs:
            paragraph.font.bold = True

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer


# --- 5. หน้าต่างการทำงาน 1: กรอกข้อมูลรายเดือน ---
if app_mode == "📝 กรอกข้อมูลรายเดือน":
    st.title("📝 บันทึกข้อมูล Plan & Action รายเดือน")
    
    col_m, col_btn = st.columns([2, 3])
    with col_m:
        selected_month = st.selectbox("📅 เลือกเดือนที่ต้องการบันทึก/แก้ไข", MONTHS)

    st.markdown(f"#### แก้ไขตัวเลขประจำเดือน : **{selected_month}**")

    # สร้าง Dataframe สำหรับแก้ไขรายเดือน
    df_current = pd.DataFrame({
        "Area / Machine": MACHINES,
        "Plan": st.session_state.plan_db[selected_month].values,
        "Action": st.session_state.action_db[selected_month].values
    })

    edited_df = st.data_editor(
        df_current,
        column_config={
            "Area / Machine": st.column_config.TextColumn("รายการเครื่องจักร", disabled=True),
            "Plan": st.column_config.NumberColumn("จำนวนแผนงาน (Plan)", min_value=0, step=1),
            "Action": st.column_config.NumberColumn("จำนวนทำจริง (Action)", min_value=0, step=1),
        },
        use_container_width=True,
        hide_index=True
    )

    # ปุ่มบันทึกข้อมูล
    if st.button("💾 บันทึกข้อมูลประจำเดือนนี้"):
        st.session_state.plan_db[selected_month] = edited_df["Plan"].values
        st.session_state.action_db[selected_month] = edited_df["Action"].values
        st.success(f"บันทึกข้อมูลเดือน {selected_month} เรียบร้อยแล้ว!")

    # คำนวณสรุปเฉพาะเดือนนี้
    m_plan = edited_df["Plan"].sum()
    m_action = edited_df["Action"].sum()
    m_kpi = (m_action / m_plan * 100) if m_plan > 0 else 0.0
    m_pass = m_kpi >= target_kpi

    st.markdown("---")
    st.subheader(f"📊 สรุปผล KPI เดือน {selected_month}")
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("ยอดรวม Plan", f"{m_plan} ครั้ง")
    c2.metric("ยอดรวม Action", f"{m_action} ครั้ง")
    c3.metric("ผลสัมฤทธิ์ KPI", f"{m_kpi:.2f}%", delta=f"{m_kpi - target_kpi:.2f}% vs Target")
    c4.metric("สถานะ", "🟢 ผ่านเกณฑ์" if m_pass else "🔴 ไม่ผ่านเกณฑ์")

    # ปุ่มสร้างและดาวน์โหลด PowerPoint รายเดือน
    ppt_file = generate_pptx_slide(selected_month, edited_df.set_index("Area / Machine")["Plan"], edited_df.set_index("Area / Machine")["Action"])
    st.download_button(
        label=f"📥 ดาวน์โหลด PowerPoint ประจำเดือน {selected_month}",
        data=ppt_file,
        file_name=f"KPI_PM_{selected_month}_{year}.pptx",
        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )

# --- 6. หน้าต่างการทำงาน 2: Dashboard ภาพรวม ---
elif app_mode == "📊 Dashboard & สรุปภาพรวม":
    st.title("📊 สรุปภาพรวม KPI ซ่อมบำรุงประจำปี")

    # คำนวณสรุปยอดรวมรายเดือน
    total_plan_series = st.session_state.plan_db.sum(axis=0)
    total_action_series = st.session_state.action_db.sum(axis=0)

    kpi_monthly = []
    for m in MONTHS:
        p = total_plan_series[m]
        a = total_action_series[m]
        kpi_val = (a / p * 100) if p > 0 else 0.0
        kpi_monthly.append(kpi_val)

    summary_df = pd.DataFrame({
        "Plan": total_plan_series,
        "Action": total_action_series,
        "KPI (%)": kpi_monthly
    })

    # คำนวณสะสมรวมทั้งปี
    year_plan = total_plan_series.sum()
    year_action = total_action_series.sum()
    year_kpi = (year_action / year_plan * 100) if year_plan > 0 else 0.0

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("แผนงานรวมทั้งปี", f"{year_plan} ครั้ง")
    c2.metric("ทำจริงรวมทั้งปี", f"{year_action} ครั้ง")
    c3.metric("KPI ภาพรวมทั้งปี", f"{year_kpi:.2f}%")
    c4.metric("สถานะรวม", "🟢 ผ่านเกณฑ์" if year_kpi >= target_kpi else "🔴 ไม่ผ่านเกณฑ์")

    st.markdown("---")
    st.subheader("📈 แนวโน้ม % KPI รายเดือน (Jan - Dec)")
    st.line_chart(summary_df["KPI (%)"])

    st.subheader("📋 ตารางข้อมูลสรุปรายเดือน (Plan vs Action vs KPI%)")
    st.dataframe(summary_df.T, use_container_width=True)
